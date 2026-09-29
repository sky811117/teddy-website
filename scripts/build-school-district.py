# -*- coding: utf-8 -*-
"""
台中市國中小學區查詢地圖 — 建置腳本
================================================================
產出（全部寫到 public/data/school-district/，給 src/pages/tools/school-district.astro 用）：
  zones_gs.json / zones_jh.json        每校一個（Multi）Polygon 學區範圍（GeoJSON FeatureCollection；地圖上色、點地圖判斷用）
                                       ⚠️ 副檔名用 .json 不用 .geojson：Cloudflare 一定會壓縮 application/json
  lookup_gs.json   / lookup_jh.json    「區|里」→ 各鄰對到哪些學校＋學區表條件片段（地址查詢用）
  school_info.json                     每校的座標、地址、電話、學區里鄰清單（點學校時顯示）
  meta.json                            學年度、來源版本、統計、官方查詢連結
搭配：public/js/tc-addr.js（地址 → 區/里/鄰）＋ public/js/sd-lookup.js（區/里/鄰 → 學校、燈號）

學年度：115 學年度（2026-08 起適用）

資料來源（全部是政府官方公告或開放資料；⛔ 不使用任何民間網站整理好的學區檔）
----------------------------------------------------------------
1. 學區表（由上游 fetch_tables.py 解析，產出 zones_*.json / by_li_*.json / li_ref.json）
   - 國小：臺中市政府教育局「115學年度本市國民小學學區表」
     https://www.central.taichung.gov.tw/media/1317817/115學年度本市國民小學學區表.pdf
     （中區區公所轉公告：https://www.central.taichung.gov.tw/831735/831749/831750/3889718）
   - 國中：臺中市政府教育局「臺中市115學年度國中學區表」
     https://www.xitun.taichung.gov.tw/media/1316223/參考附件2臺中市115學年度國中學區表.odt
     （西屯區公所公告：https://www.xitun.taichung.gov.tw/831932/832178/832214/3885985）
2. 村里界、學校座標（由上游 build_geo.py 產出 villages_tc.geojson / schools.json）
   - 內政部國土測繪中心「村里界圖(TWD97經緯度)」VILLAGE_NLSC_1150817  https://data.gov.tw/dataset/7438
   - 教育部統計處 115 學年度各級學校名錄（校址、電話）https://stats.moe.gov.tw/files/school/115/（e1_new / j1_new / aj_new / ac）
   - 教育部統計處「115學年度國中小異動一覽表」（新設校）https://stats.moe.gov.tw/files/school/115/115_BasicChange.xlsx
3. 門牌點位（每個門牌的 區/里/鄰/座標；切學區多邊形、檢查涵蓋率）
   - 讀網站自己的門牌庫 public/data/tc-addr/（scripts/build-tc-addr.py 產生；
     來源：臺中市政府資料開放平臺「臺中市空間資訊建物及門牌號碼位置新版本資料」115年1月 GIS 門牌）
     → 地圖用的里鄰跟網頁地址查詢用的是同一份，兩邊不會各說各話

怎麼更新（教育局公告新學年度學區表時，通常每年 1–3 月）
----------------------------------------------------------------
1. 重跑上游學區表解析：python <scratch>/tools2/sd_tables/fetch_tables.py --refresh
   （新年度要先把 fetch_tables.py 裡的 PDF / ODT 網址換成新公告的；解析後看 report.json 的 issues）
2. 村里界或學校名錄有新版時：python <scratch>/tools2/geo/build_geo.py --refresh
   （新學年度也要把下面 MOE_YEAR、BASIC_CHANGE_URL 換掉）
3. 門牌換新月份時：先照 scripts/build-tc-addr.py 檔頭步驟更新 public/data/tc-addr/
4. 執行本腳本：
     python scripts/build-school-district.py [--sd <sd_tables資料夾>] [--geo <geo資料夾>] [--addr <門牌庫資料夾>]
   （預設路徑見下面 DEFAULT_*；也可用環境變數 SD_TABLES_DIR / SD_GEO_DIR）
5. 腳本最後會跑一致性檢查，任何一項不過就 exit 1、不會留下半套檔案（先寫暫存檔、全部通過才換上）。
   其中「路段條件」兩項（check_cond_fragments / check_lis_conds）是拿學區表原文另外推一次每校每鄰應有的片段來對：
   原文本校在某鄰「不分路段」也有、另外某路段一側是共同學區 → 查詢表一定要同時有不分路段和帶條件兩段；
   原文本校只有帶條件的部分 → 不能多出不分路段；學校彈窗描述的條件要掛在正確的鄰上。
   原文就沒有不分路段部分的鄰列在報告 gs_cond_check.cond_only / jh_cond_check.cond_only。
6. 本機用 http://localhost:4321/tools/school-district/ 驗幾個地址，確認再部署。

學區多邊形怎麼畫（zones_*.json）
----------------------------------------------------------------
- 全市每個門牌點依「它的里鄰對到的學校組合」上標籤，整個臺中市做一次 Voronoi（每個門牌點的勢力範圍），
  同標籤合併、裁到村里界的市界內，再用 coverage_simplify（相鄰區塊共用邊，不會出現縫或重疊）把鋸齒修順。
  → 點地圖任一處＝離那裡最近的門牌所屬的學區，跟地址查詢用同一份門牌資料；交界是依門牌位置推估的。
  （舊版用村里界＋里內切鄰，村里界跟門牌資料的里別有出入的地方會畫錯，全市約 0.8% 門牌落在別校範圍；
    改成全市門牌 Voronoi 後約 0.01%，只剩交界修順的誤差）
- 共同學區／依路段拆分（同一鄰對到兩校以上）那塊會同時出現在每一校的多邊形裡（重疊），properties.shared / split 列出。
- 整區學區（中科實中國小部＝大雅區＋西屯區共同學區；和平國中＝和平區全區自由學區）properties.wide = true，
  用村里界拼出整區；頁面只畫外框、不填色，免得蓋住一般學區。

輸出格式（全部 UTF-8、緊湊 JSON；學校一律用「校 id」＝學區表校名，同名校加區名，例「成功國小(東區)」）
----------------------------------------------------------------
zones_gs.json / zones_jh.json：FeatureCollection，每校一個 Feature（Polygon / MultiPolygon，WGS84 5 位小數）
  properties = {school: 校id, name: 顯示校名, level: '國小'|'國中', dist: 學校所在區, n_li: 學區涉及幾個里,
                shared: [學區表寫明共同學區的其他校id], split?: [同一鄰依路段／條件拆給的其他校id],
                approx: 學區裡有依鄰劃分的里, ci?: 配色序號 0–4（相鄰學區不同號）, free?: 有自由學區,
                wide?: 整區學區, wide_dists?: [區], wide_raw?: 學區表原文}
lookup_gs.json / lookup_jh.json：
  {"區|里": {all:[校id], lin:{"鄰":[校id]}, free?:{校id:"all"|[鄰]},
             fr?:{"鄰"|"all": [{s:[校id], c:"條件文字", co:1?, f:[自由學區校id]?}]}},
   "_dist_wide": {區:[{school, free_zone, co_zone, raw, info?:{text, url}}]}}
  查某里某鄰 = all ∪ lin[鄰]（再加 _dist_wide[區] 當補充說明）；625 個里都有 key（沒學校的是空陣列）。
  fr＝學區表逐筆原文拆出的「片段」：同一鄰在學區表出現好幾次（例：仁和里第 18 鄰「崇德八路以南」只屬松竹國小、
  「崇德八路以北」為松竹、松強共同學區）就有好幾個片段；co=1 表示原文寫明「共同學區」。
  本校整鄰都是學區、只有某路段一側另外是共同學區（建功國小「水景里(全里)…【9鄰(74號快速道路以東)】為廍子國小及建功國小共同學區」）
  → 那一鄰一定同時有 {s:[本校], c:''}（不分路段）和 {s:[本校, 別校], c:'74號快速道路以東', co:1} 兩段。
  只有單一學校、沒條件、不是共同／自由學區的鄰不列 fr（網頁端當成一個沒條件的片段）。
school_info.json：{gs:{校id:{name, level, dist, addr, lat, lng, phone, lis:[[區, 里, 鄰描述, 共同學區校id…]],
                   raw:[學區表原文…], status?, status_src?, geo?, notes?}}, jh:{…}}
  lis 的「共同學區校id」只掛在學區表原文寫明共同學區的那一列。
meta.json：{title, year, yearNote, built, counts, officialQuery:{name,url}, tables:[學區表公告],
            sources:[{name, url, page?, version, kind:'opendata'|'official'|'other', org?, year?}],
            license, attribution:[顯名聲明], coverage:{gs:{…}, jh:{…}}, notes:[…]}
"""
import argparse
import collections
import datetime
import glob
import json
import os
import re
import sys
import time

import numpy as np
import shapely
from shapely.geometry import mapping, shape
from pyproj import Transformer

sys.stdout.reconfigure(encoding='utf-8')

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(REPO, 'public', 'data', 'school-district')
_SCRATCH = r'C:\Users\a0920\AppData\Local\Temp\claude\C--Users-a0920--claude\f8c88a94-377a-48d9-b470-cd09a6cc09e8\scratchpad\tools2'
DEFAULT_SD = os.environ.get('SD_TABLES_DIR', os.path.join(_SCRATCH, 'sd_tables'))
DEFAULT_GEO = os.environ.get('SD_GEO_DIR', os.path.join(_SCRATCH, 'geo'))
DEFAULT_ADDR = os.path.join(REPO, 'public', 'data', 'tc-addr')
DEFAULT_REPORT = os.path.join(_SCRATCH, 'build_school', 'build_report.json')

YEAR = '115學年度'
MOE_YEAR = '115'
BASIC_CHANGE_URL = f'https://stats.moe.gov.tw/files/school/{MOE_YEAR}/115_BasicChange.xlsx'
MOE_LIST_URL = f'https://stats.moe.gov.tw/files/school/{MOE_YEAR}/'
OFFICIAL_QUERY = {
    'name': '臺中市政府教育局 學區查詢',
    'url': 'https://www.tc.edu.tw/page/02b0fa2f-7dda-404f-b411-8286cd97c9c1',
}
# 整區學區的補充說明（頁面照原文顯示，不自己推論「可以選讀」）
WIDE_INFO = {
    '國立中科實驗高級中學國小部': {
        'text': '國小部自己辦理招生（網路報名、抽籤）',   # 頁面接「，名額和設籍規定：看學校招生資訊」連結
        'url': 'https://www.nehs.tc.edu.tw/招生資訊國小部/',   # 2026-09-29 讀過：列有 115 學年度國小部招生簡章、抽籤結果公告
    },
}

DIST_BY_CODE = {
    '6600100': '中區', '6600200': '東區', '6600300': '南區', '6600400': '西區', '6600500': '北區',
    '6600600': '西屯區', '6600700': '南屯區', '6600800': '北屯區', '6600900': '豐原區', '6601000': '東勢區',
    '6601100': '大甲區', '6601200': '清水區', '6601300': '沙鹿區', '6601400': '梧棲區', '6601500': '后里區',
    '6601600': '神岡區', '6601700': '潭子區', '6601800': '大雅區', '6601900': '新社區', '6602000': '石岡區',
    '6602100': '外埔區', '6602200': '大安區', '6602300': '烏日區', '6602400': '大肚區', '6602500': '龍井區',
    '6602600': '霧峰區', '6602700': '太平區', '6602800': '大里區', '6602900': '和平區',
}
DISTS = list(DIST_BY_CODE.values())
DISTS_BY_LEN = sorted(DISTS, key=len, reverse=True)
LEVELS = [('gs', '國小'), ('jh', '國中')]

# 里名異體字（學區表 / 舊資料寫法 → 村里界與門牌庫寫法）
LI_VARIANTS = [('臺', '台'), ('殼', '壳'), ('館', '舘'), ('雙', '双'), ('壩', '埧'), ('褔', '福'), ('豊', '豐')]
LI_ALIAS = {'神州里': '神洲里', '中科里': '中嵙里'}
# 官方學區表原文的簡體字（國中學區表把「后綜高中」的綜打成簡體 U+7EFC）→ 校 id、校名、原文一律改回繁體
SIMP_FIX = {'\u7efc': '綜'}

# 籌備中 / 預定設校：不列入（校名或狀態含這些字就排除，並寫進報告）
NOT_OPEN_RE = re.compile(r'籌備|預定|擬設|規劃中')

VORONOI_SMOOTH_M = 10.0     # 學區交界修順容差（公尺）；10m 時門牌點落到別校範圍約 0.01%
MIN_HOLE_M2 = 40.0          # 小於這個面積的洞（多半是切割誤差）補起來
MIN_PART_M2 = 15.0          # 小於這個面積的碎片丟掉

T_FWD = Transformer.from_crs(4326, 3826, always_xy=True)   # WGS84 → TWD97 TM2（公尺）
T_INV = Transformer.from_crs(3826, 4326, always_xy=True)


def to_m(g):
    return shapely.transform(g, lambda xy: np.column_stack(T_FWD.transform(xy[:, 0], xy[:, 1])))


def to_deg(g):
    return shapely.transform(g, lambda xy: np.column_stack(T_INV.transform(xy[:, 0], xy[:, 1])))


def jdump(obj, path):
    with open(path, 'w', encoding='utf-8', newline='\n') as fh:
        json.dump(obj, fh, ensure_ascii=False, separators=(',', ':'))


def jload(path):
    with open(path, encoding='utf-8') as fh:
        return json.load(fh)


def log(*a):
    print(*a, flush=True)


def fix_simp(obj):
    """整份結構裡的字串（含 dict key）把簡體錯字改回繁體。"""
    if isinstance(obj, str):
        for a, b in SIMP_FIX.items():
            obj = obj.replace(a, b)
        return obj
    if isinstance(obj, list):
        return [fix_simp(x) for x in obj]
    if isinstance(obj, dict):
        return {fix_simp(k): fix_simp(v) for k, v in obj.items()}
    return obj


# ─────────────────────────── 載入 ───────────────────────────
def load_inputs(sd, geo):
    need = {
        'zones_gs': os.path.join(sd, 'zones_gs.json'), 'zones_jh': os.path.join(sd, 'zones_jh.json'),
        'by_li_gs': os.path.join(sd, 'by_li_gs.json'), 'by_li_jh': os.path.join(sd, 'by_li_jh.json'),
        'li_ref': os.path.join(sd, 'li_ref.json'), 'sd_report': os.path.join(sd, 'report.json'),
        'villages': os.path.join(geo, 'villages_tc.geojson'), 'schools': os.path.join(geo, 'schools.json'),
        'max_lin': os.path.join(geo, 'max_lin.json'), 'geo_report': os.path.join(geo, 'geo_report.json'),
    }
    miss = [p for p in need.values() if not os.path.exists(p)]
    if miss:
        raise SystemExit('找不到上游檔案：\n  ' + '\n  '.join(miss))
    out = {k: jload(p) for k, p in need.items()}
    for k in ('zones_gs', 'zones_jh', 'by_li_gs', 'by_li_jh'):
        out[k] = fix_simp(out[k])
    return out


def load_moe_phones(geo):
    """教育部 115 學年度學校名錄 → {學校代碼: 電話}（只取臺中市）；讀不到就回空（沿用學區表的電話）。"""
    phones = {}
    try:
        import openpyxl
    except ImportError:
        return phones
    for fn in ('e1_new.xlsx', 'j1_new.xlsx', 'aj_new.xlsx', 'ac.xlsx', '115_BasicChange.xlsx'):
        p = os.path.join(geo, 'raw', 'moe115', fn)
        if not os.path.exists(p):
            continue
        wb = openpyxl.load_workbook(p, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(x).strip() for x in row if x is not None]
                if len(cells) < 3 or not re.fullmatch(r'\d{6}', cells[0]):
                    continue
                if not any('臺中市' in c for c in cells):
                    continue
                ph = next((c for c in cells if re.match(r'^\(0\d{1,3}\)\d', c)), None)
                if ph:
                    phones.setdefault(cells[0], ph)
    return phones


def norm_phone(ph):
    """電話統一成顯示用字串：學區表 '22124224' → '04-22124224'；'049-2523020' 不動；
    教育部名錄 '(04)24360166#826' → '04-24360166#826'、'(049)2523020' → '049-2523020'；'未定'、空白 → ''。"""
    ph = (ph or '').strip().replace(' ', '')
    m = re.match(r'^\((0\d{1,3})\)(\d.*)$', ph)
    if m:
        return f'{m.group(1)}-{m.group(2)}'
    if re.match(r'^0\d{1,3}-\d{6,8}', ph):
        return ph
    if re.match(r'^\d{7,8}(#\d+)?$', ph):
        return '04-' + ph
    return ''


def load_new_schools(geo):
    """教育部「115學年度國中小異動一覽表」的「※新設」段 → {學校代碼}。"""
    out = set()
    p = os.path.join(geo, 'raw', 'moe115', '115_BasicChange.xlsx')
    if not os.path.exists(p):
        return out
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True)
    sec = ''
    for row in wb.worksheets[0].iter_rows(values_only=True):
        cells = [str(x).strip() for x in row if x is not None]
        if not cells:
            continue
        if cells[0].startswith('※'):
            sec = cells[0]
            continue
        if sec == '※新設' and re.fullmatch(r'\d{6}', cells[0]):
            out.add(cells[0])
    return out


def load_addr_points(addr_dir):
    """門牌庫（public/data/tc-addr）→ 每個門牌的 (經度, 緯度, 區, 里, 鄰)；
    cnt：(區, 里, 鄰) → 門牌數（主要里鄰）；alt_cnt：同一門牌也登記在別的里鄰（分片 x）的次數。"""
    ix = jload(os.path.join(addr_dir, 'index.json'))
    o_lat, o_lng, sc = ix['o']
    dn = [n for _, n in ix['d']]
    pts = []
    cnt, alt_cnt = collections.Counter(), collections.Counter()
    files = sorted(glob.glob(os.path.join(addr_dir, 's', '*.json')))
    if not files:
        raise SystemExit(f'門牌庫是空的：{addr_dir}（先跑 scripts/build-tc-addr.py）')
    for fn in files:
        sh = jload(fn)
        di = sh['d']
        dist, lis = dn[di], ix['li'][di]
        for ro in sh['r'].values():
            xs = ro.get('x') or {}
            for gk, gs in ro['g'].items():
                alts = {}
                if xs.get(gk):
                    for part in xs[gk].split(';'):
                        no, rest = part.split('=', 1)
                        alts[no] = [tuple(int(v) for v in a.split(':')) for a in rest.split('|')[1:]]
                li = lin = la = ln = None
                for part in gs.split(';'):
                    no, f_li, f_lin, f_la, f_ln = part.split(',')
                    if f_li != '':
                        li = int(f_li)
                    if f_lin != '':
                        lin = int(f_lin)
                    if la is None:
                        la, ln = int(f_la), int(f_ln)
                    else:
                        la += int(f_la)
                        ln += int(f_ln)
                    pts.append((ln / sc + o_lng, la / sc + o_lat, dist, lis[li], lin))
                    cnt[(dist, lis[li], lin)] += 1
                    for a_li, a_lin, _c in alts.get(no, []):
                        alt_cnt[(dist, lis[a_li], a_lin)] += 1
    return pts, cnt, alt_cnt, ix.get('ver'), ix.get('n')


# ─────────────────────────── 學校對座標 ───────────────────────────
def dist_of_addr(addr):
    a = re.sub(r'^(臺中市|台中市)', '', addr or '')
    for d in DISTS_BY_LEN:
        if a.startswith(d):
            return d
    return None


def match_schools(level, zones, geo_schools, li_dists):
    """學區表的學校 → 教育部名錄＋門牌定位的學校點。回傳 {id: geo_rec}、錯誤清單。"""
    idx = collections.defaultdict(list)
    for s in geo_schools:
        if s['level'] != level:
            continue
        for n in [s['short']] + s.get('aliases', []):
            idx[fix_simp(n)].append(s)
    out, errs = {}, []
    for z in zones:
        name = z['school']
        cands = list({id(c): c for c in idx.get(name, [])}.values())
        d = dist_of_addr(z.get('addr'))
        if not d:   # 校址沒寫區（例：和平國小「南勢里東關路三段54號」）→ 用里名推
            m = re.match(r'^(.{1,3}?[里村])', z.get('addr') or '')
            if m and len(li_dists.get(m.group(1), ())) == 1:
                d = next(iter(li_dists[m.group(1)]))
        exact = [c for c in cands if c['short'] == name]   # 簡稱完全相同的優先於別名
        if exact:
            cands = exact
        if len(cands) > 1:
            f = [c for c in cands if c['dist'] == (d or z.get('home_dist'))]
            cands = f
        if len(cands) != 1:
            errs.append(f'{level} {z["id"]}：名錄對到 {len(cands)} 筆（校址區 {d}）')
            continue
        if d and cands[0]['dist'] != d:
            errs.append(f'{level} {z["id"]}：學區表校址在{d}，名錄對到的是{cands[0]["dist"]}{cands[0]["short"]}')
            continue
        out[z['id']] = cands[0]
    return out, errs


# ─────────────────────────── 查詢表 ───────────────────────────
def li_candidates(li):
    c = [li, LI_ALIAS.get(li, li)]
    s = li
    for a, b in LI_VARIANTS:
        s = s.replace(a, b)
    c.append(s)
    c.append(LI_ALIAS.get(s, s))
    return list(dict.fromkeys(c))


def check_range_villages(level, zones, vill_keys, li_dists):
    """學區表每一筆範圍的里都要在村里界找得到；對不上時依序試：異體字對照表 → 全市唯一同名里。"""
    fixed, bad = [], []
    for z in zones:
        for r in z['ranges']:
            if r.get('whole_dist') or r['li'] == '*':
                if r['dist'] not in DISTS:
                    bad.append(f'{z["id"]} 整區範圍的區名「{r["dist"]}」不存在')
                continue
            li = r.get('li_norm') or r['li']
            if (r['dist'], li) in vill_keys:
                if r['li'] != li:
                    fixed.append(f'{level} {z["id"]}：{r["dist"]}{r["li"]} → {r["dist"]}{li}（異體字對照）')
                continue
            hit = next(((r['dist'], c) for c in li_candidates(li) if (r['dist'], c) in vill_keys), None)
            if not hit:
                for c in li_candidates(li):
                    if len(li_dists.get(c, ())) == 1:
                        hit = (next(iter(li_dists[c])), c)
                        break
            if hit:
                fixed.append(f'{level} {z["id"]}：{r["dist"]}{r["li"]} → {hit[0]}{hit[1]}')
            else:
                bad.append(f'{level} {z["id"]}：{r["dist"]}{r["li"]}（原文 {r.get("raw", "")[:40]}）')
    return fixed, bad


# ─── 共同學區的補強（學區表原文寫明、但上游解析沒標到 co 的情況）───
# 1. 範圍自己的註記只講部分鄰：「和平里(全里)【第14鄰(軍福十三路以南)、16、33…鄰為廍子國中共同學區】」
#    → 共同學區只套在列出的那幾鄰（有括號條件的鄰：共同學區只限那個條件，其餘仍是本校）
# 2. 註記點名了別校：「土庫里(第3-9…鄰)【為崇倫、向上國中共同學區】」而向上國中那一列沒寫
#    → 同一鄰、同條件的向上國中片段併進這個共同學區片段
# 3. 學校註記（註記原文／備註欄）寫「與國安國小共同學區」「中社里15、39-42為大秀國小、槺榔國小之共同學區」
#    → 同一鄰只有這幾校各自一段、條件相同（學區表重複列）時，併成共同學區
CO_PARTNER_RE = re.compile(r'(?:為|與)([^為與【】※；]*?)(?:之|的)?共同學區')
CO_LINS_RE = re.compile(r'【(?:其中)?(第[^【】]*?)為[^【】]*?共同學區】')
NOTE_PHRASE_RE = re.compile(r'([^／：:；;]*?)(?:之|的)?共同學區')
NOTE_LIST_RE = re.compile(r'([^\s:：;；,，、／]{1,5}里)((?:\d+(?:-\d+)?、?)+)為([^；;／]+?)之?共同學區')
SCHOOL_SUFFIX_RE = re.compile(r'(國民小學|國民中學|高級中學國小部|高級中學國中部|高中國中部|高中國小部|國中小|國小部|國中部|國小|國中|高中)$')


def school_short(sid):
    n = re.sub(r'\(.*?\)', '', sid)
    n = re.sub(r'^(國立|市立|私立)', '', n)
    n = re.sub(r'實驗高級中學.*$', '', n)
    n = SCHOOL_SUFFIX_RE.sub('', n)
    return n


def parse_lin_tokens(txt):
    """「第3-6、12、17鄰及14鄰(大慶街二段以南)」→ {3:'',4:'',5:'',6:'',12:'',17:'',14:'大慶街二段以南'}；看不懂回 None。"""
    out = {}
    for tok in re.split(r'[、及]', txt):
        tok = tok.strip()
        if not tok:
            continue
        m = re.fullmatch(r'第?(\d+)(?:\s*[-~～至]\s*(\d+))?鄰?\s*(?:[(（]([^)）]*)[)）])?鄰?', tok)
        if not m:
            return None
        a, b, c = int(m.group(1)), int(m.group(2) or m.group(1)), (m.group(3) or '').strip()
        if b < a or b - a > 60:
            return None
        for n in range(a, b + 1):
            out[n] = c
    return out or None


ANN_CO_RE = re.compile(r'【[^【】]*共同學區[^【】]*】')


def base_partial(r):
    """範圍「本校自己」的鄰別條件 {鄰(str): 條件}。
    上游把【…共同學區】註記裡的路段條件也放進 partial：
      四育國中「樹德里(第1-6、11-33鄰)【其中第3-6…鄰及14鄰(大慶街二段以南)為光德國中共同學區】」→ partial {14: 大慶街二段以南}
    但原文的意思是四育國中整個 14 鄰都是學區，只有大慶街二段以南「另外」是和光德國中的共同學區
    → 條件文字只出現在註記裡、註記外面又有本校自己的鄰別時，這個條件不算本校的條件（留給共同學區那一段）。
    整筆就是註記（光德國中「樹德里【第3-6…及14鄰(大慶街二段以南)為四育國中共同學區】」）時照舊。"""
    part = {str(a): (b or '').strip() for a, b in (r.get('partial') or {}).items()}
    if not part or not r.get('co_zone'):
        return part
    raw = r.get('raw') or ''
    anns = ANN_CO_RE.findall(raw)
    rest = ANN_CO_RE.sub('', raw)
    if not anns or not re.search(r'\d|全里', rest):
        return part
    return {L: c for L, c in part.items() if c in rest or not any(c in a for a in anns)}


def co_lins_of(r):
    """共同學區只限部分鄰時回 {鄰: 條件}；整筆都是共同學區（或看不懂）回 None。"""
    if not r.get('co_zone'):
        return None
    txt = (r.get('ann') or '') + '；' + (r.get('raw') or '')
    anns = re.findall(r'【[^【】]*共同學區[^【】]*】', txt)
    if not anns:
        return None
    out = {}
    for a in anns:
        m = CO_LINS_RE.fullmatch(a)
        if not m:
            return None           # 有一段是「整筆共同學區」→ 不限鄰
        d = parse_lin_tokens(m.group(1))
        if d is None:
            return None
        for n, c in d.items():
            out[n] = c
    return out


def build_fragments(zones, fixlog=None):
    """學區表逐筆範圍 → {'區|里': {'鄰'|'all': [片段]}}，片段 = {s:[校id], c:條件文字, co:bool, f:[自由學區校id]}。
    同一鄰在學區表出現好幾次（單獨的一段＋共同學區的一段、依路段分給兩校…）就有好幾個片段。"""
    fixlog = fixlog if fixlog is not None else []
    note_merged = collections.defaultdict(set)     # (校id, 區|里) → 依學校註記併成共同學區的鄰
    shorts = {z['id']: school_short(z['id']) for z in zones}
    raw_fr = collections.defaultdict(lambda: collections.defaultdict(dict))

    def put(k, L, gid, sid, c, co, fz_hit, ann, q=None):
        f = raw_fr[k][L].setdefault(gid, {'s': [], 'c': c, 'co': co, 'f': [], 'ann': [], 'q': {}})
        if sid not in f['s']:
            f['s'].append(sid)
        if q and co:
            f['q'][sid] = q            # 「可選…就讀」註記原文（原文沒寫共同學區；見 main 的 choice）
        if fz_hit and sid not in f['f']:
            f['f'].append(sid)
        if ann and ann not in f['ann']:
            f['ann'].append(ann)

    whole_override = collections.defaultdict(set)   # (區|里, 校id) → 全里範圍另外寫了條件的鄰（這些鄰不算全里的「不分路段」）
    for z in zones:
        sid = z['id']
        for r in z['ranges']:
            if r.get('whole_dist') or r['li'] == '*':
                continue
            k = f'{r["dist"]}|{r.get("li_norm") or r["li"]}'
            lins = ['all'] if r['lin'] == 'all' else [str(n) for n in r['lin']]
            part = base_partial(r)          # 只取本校自己的鄰別條件（註記裡共同學區那一段的條件另外放）
            fz = r.get('free_zone')
            co = bool(r.get('co_zone'))
            cl = co_lins_of(r)
            ann = (r.get('ann') or '') + (r.get('raw') or '') if co else ''
            if cl is not None:
                fixlog.append(f'共同學區只限部分鄰：{sid} {k} 第{sorted(cl)}鄰')
            if r['lin'] == 'all' and part:
                # 全里範圍另外寫某一鄰的條件（旭光國小「學田里(全里)(12鄰:312巷105弄以外)」）：
                # 那一鄰本校只有這個條件，另放一段；全里那段不能套到這一鄰（見下面「全里不分路段補到帶條件的鄰」）
                lins = lins + sorted(part, key=int)
                whole_override[(k, sid)] |= set(part)
            for L in lins:
                c = (part.get(L) or r.get('cond') or '').strip() if L != 'all' else (r.get('cond') or '').strip()
                fz_hit = fz is True or (isinstance(fz, list) and L != 'all' and int(L) in fz)
                if co and cl is None:
                    put(k, L, ('co', r.get('raw', ''), c), sid, c, True, fz_hit, ann, r.get('choice'))
                    continue
                if co and L != 'all' and int(L) in cl and (not cl[int(L)] or cl[int(L)] == c):
                    put(k, L, ('co', r.get('raw', ''), c), sid, c, True, fz_hit, ann, r.get('choice'))
                    continue
                put(k, L, ('solo', sid, c), sid, c, False, fz_hit, '')
            if cl is not None:
                # 共同學區的鄰：沒條件的上面已放；有條件的（例：14鄰(軍福十三路以南)）另放一段「條件＋共同學區」；
                # 範圍是全里時，列出的鄰也要各放一段（全里那段在網頁端會併進每一鄰）
                for n, cc in sorted(cl.items()):
                    if r['lin'] == 'all':
                        base = (part.get(str(n)) or r.get('cond') or '').strip()
                        c = cc or base
                        if cc and cc != base:
                            # 例：和平里第 14 鄰「軍福十三路以南」才是共同學區 → 這一鄰另放本校不分路段的一段，網頁才會列「依路段分」
                            put(k, str(n), ('solo', sid, base), sid, base, False, False, '')
                    else:
                        cL = (part.get(str(n)) or r.get('cond') or '').strip()
                        if n not in r['lin'] or not cc or cc == cL:
                            continue          # 上面已經放成共同學區
                        c = cc
                    put(k, str(n), ('co', r.get('raw', ''), c), sid, c, True, False, ann, r.get('choice'))

    # 全里不分路段補到帶條件的鄰：
    #   建功國小「水景里(全里)」＋「水景里【9鄰(74號快速道路以東)、34鄰(…)】為廍子國小及建功國小共同學區」
    #   → 9 鄰只有「74號快速道路以東：廍子、建功（共同學區）」一段的話，網頁端看到建功已經有片段，就不會再補「不分路段」，
    #     路另一側（只屬建功）會被說成共同學區。所以這一鄰要明寫一段「建功國小 不分路段」（跟國中和平里 14 鄰同一個做法）。
    #   全里那段是共同學區（網頁端 fr.all 會併進每一鄰）或本校在這一鄰另有條件（whole_override）的不補。
    for k, by_lin in raw_fr.items():
        whole = {}
        for f in (by_lin.get('all') or {}).values():
            if not f['c'] and not f['co']:
                for s in f['s']:
                    whole.setdefault(s, f)
        if not whole:
            continue
        for L in sorted((x for x in by_lin if x != 'all'), key=int):
            for sid, fa in whole.items():
                if L in whole_override.get((k, sid), ()):
                    continue
                mine = [f for f in by_lin[L].values() if sid in f['s']]
                if any(f['c'] for f in mine) and not any(not f['c'] for f in mine):
                    put(k, L, ('solo', sid, ''), sid, '', False, sid in fa['f'], '')
                    fixlog.append(f'全里不分路段補到帶條件的鄰：{sid} {k} 第{L}鄰（原本只有 {[f["c"] for f in mine]}）')

    # 學校註記裡的共同學區說法（見上面第 3 點）
    note_pairs = []           # [(學校集合, 限定的(里, 鄰集合) 或 None, 說明)]
    for z in zones:
        notes = z.get('notes') or ''
        for m in NOTE_LIST_RE.finditer(notes):
            lins = parse_lin_tokens(m.group(2).rstrip('、'))
            ss = {sid for sid, sh in shorts.items() if len(sh) >= 2 and sh in m.group(3)}
            if lins and len(ss) >= 2:
                note_pairs.append((ss, (m.group(1), set(lins)), m.group(0)))
        for m in NOTE_PHRASE_RE.finditer(notes):
            if NOTE_LIST_RE.search(m.group(0)):
                continue
            ph = m.group(1)
            ss = {sid for sid, sh in shorts.items() if len(sh) >= 2 and sh in ph} | {z['id']}
            if len(ss) >= 2:
                note_pairs.append((ss, None, m.group(0)))

    out = {}
    for k, by_lin in raw_fr.items():
        li = k.split('|', 1)[1]
        for L, frs in by_lin.items():
            lst = list(frs.values())
            # 2. 共同學區片段的註記點名的學校：同一鄰、同條件的單獨片段併進來
            for F in [f for f in lst if f['co']]:
                ptxt = '、'.join(g for a in F['ann'] for g in CO_PARTNER_RE.findall(a))
                if not ptxt:
                    continue
                for G in [g for g in lst if not g['co'] and g['c'] == F['c']]:
                    if G in lst and all(len(shorts.get(x, '')) >= 2 and shorts.get(x, '') in ptxt and x not in F['s'] for x in G['s']):
                        F['s'] += G['s']
                        F['f'] += [x for x in G['f'] if x not in F['f']]
                        lst.remove(G)
                        fixlog.append(f'共同學區補上點名的學校：{k} 第{L}鄰 {F["s"]}（{ptxt}）')
            # 3. 學校註記：同一鄰、同一條件，這幾校各自一段（學區表重複列）→ 併成共同學區。
            #    註記沒寫是哪個里鄰的（「與國安國小共同學區」）：這一鄰只能有這幾段；
            #    註記寫明里鄰的（「中社里15、39-42為大秀國小、槺榔國小之共同學區」）：同一鄰還有別校的路段也照併
            by_c = collections.defaultdict(list)
            for g in lst:
                if not g['co'] and len(g['s']) == 1:
                    by_c[g['c']].append(g)
            for c0, grp in by_c.items():
                if len(grp) < 2:
                    continue
                ids = {g['s'][0] for g in grp}
                for ss, where, txt in note_pairs:
                    if not ids <= ss:
                        continue
                    if where is None and len(lst) != len(grp):
                        continue
                    if where and (where[0] != li or L == 'all' or int(L) not in where[1]):
                        continue
                    lst = [g for g in lst if g not in grp] + [{'s': sorted(ids), 'c': c0, 'co': True,
                                                               'f': sorted({x for g in grp for x in g['f']}), 'ann': [txt]}]
                    fixlog.append(f'學校註記標共同學區：{k} 第{L}鄰 {sorted(ids)}（{txt}）')
                    for x in ids:
                        note_merged[(x, k)].add(L)
                    break
            # 同一鄰、同一條件的「共同學區」片段，各校那一列原文寫法不同而沒併到的 → 併成一段
            co1 = [f for f in lst if f['co']]
            merged = {}
            for f in co1:
                mk = f['c']
                if mk in merged:
                    for x in f['s']:
                        if x not in merged[mk]['s']:
                            merged[mk]['s'].append(x)
                    merged[mk]['f'] += [x for x in f['f'] if x not in merged[mk]['f']]
                    merged[mk]['q'].update(f.get('q') or {})
                else:
                    merged[mk] = {'s': list(f['s']), 'c': f['c'], 'co': True, 'f': list(f['f']), 'q': dict(f.get('q') or {})}
            lst = [f for f in lst if not f['co']] + list(merged.values())
            out.setdefault(k, {})[L] = lst
    # 依學校註記併成共同學區的鄰：該校那一列（鄰全部在內）也標成共同學區，學校彈窗的學區清單才會寫出共同學區的學校
    for z in zones:
        for r in z['ranges']:
            if r.get('co_zone') or r.get('whole_dist') or r['li'] == '*' or r['lin'] == 'all':
                continue
            got = note_merged.get((z['id'], f'{r["dist"]}|{r.get("li_norm") or r["li"]}'))
            if got and {str(n) for n in r['lin']} <= got:
                r['co_zone'] = True
                fixlog.append(f'學區表列同時標共同學區：{z["id"]} {r["dist"]}{r["li"]} {r.get("raw", "")}')
    return out


def frag_json(f):
    o = {'s': sorted(f['s']), 'c': f['c']}
    if f['co']:
        o['co'] = 1
    if f['f']:
        o['f'] = sorted(f['f'])
    if f['co'] and f.get('q'):
        o['q'] = {k: f['q'][k] for k in sorted(f['q'])}   # 「可選…就讀」註記原文（校id → 原文）
    return o


def lin_school_set(e, L):
    """某里某鄰（L 是鄰號字串）對到的全部學校：all ∪ lin[L] ∪ 片段（fr[L]＋fr.all）裡的學校。"""
    fr = e.get('fr') or {}
    out = set(e.get('all') or []) | set((e.get('lin') or {}).get(L, []))
    for f in fr.get(L, []) + fr.get('all', []):
        out |= set(f['s'])
    return out


def clean_single_co(lk):
    """只有一校的鄰，片段不標共同學區（co）：上游把整列的【共同學區】註記套到那一列所有鄰
    （例：美群國小「瑞城里(1-3 鄰、4鄰【共同學區】18-21鄰)」→ 1-3、18-21 鄰也被標 co，但這幾鄰只有美群國小）。
    網頁端只有一校時本來就不看 co（燈號不受影響），這裡把資料弄乾淨，免得以後改 classify 時被誤判成共同學區。
    「全里」(all) 的片段會套到每一鄰，其中有的鄰真的有別校共同 → 不能拿掉，只列進報告
    （sd-lookup.js 的 classify 一定要先看學校數：只有一校就是 single/cond，不看 co）。
    回傳 (清掉的 [區|里 第N鄰], 全里片段標 co 但有些鄰只有一校的 [區|里])。"""
    cleaned, all_single = [], []
    for k, e in lk.items():
        if k.startswith('_') or not e.get('fr'):
            continue
        fr = e['fr']
        for L in [x for x in fr if x != 'all']:
            if len(lin_school_set(e, L)) != 1 or not any(f.get('co') for f in fr[L]):
                continue
            for f in fr[L]:
                f.pop('co', None)
                f.pop('q', None)
            cleaned.append(f'{k} 第{L}鄰')
            if len(fr[L]) == 1 and not fr[L][0].get('c') and not fr[L][0].get('f'):
                del fr[L]            # 只剩一個沒條件的片段 → 跟沒列一樣
        if any(f.get('co') for f in fr.get('all', [])):
            lins = set(e.get('lin') or {}) | {x for x in fr if x != 'all'}
            if any(len(lin_school_set(e, L)) == 1 for L in lins | {'-1'}):
                all_single.append(k)
        if not fr:
            del e['fr']
    return cleaned, all_single


def check_single_co(lk):
    """一致性檢查：逐鄰列的片段（不含全里片段），鄰裡只有一校就不能標共同學區。"""
    bad = []
    for k, e in lk.items():
        if k.startswith('_'):
            continue
        for L, lst in (e.get('fr') or {}).items():
            if L != 'all' and any(f.get('co') for f in lst) and len(lin_school_set(e, L)) == 1:
                bad.append(f'{k} 第{L}鄰只有一校，片段卻標了共同學區')
    return bad


def build_lookup(by_li, frags, zones, vill_keys, school_ids, excluded):
    """上游 by_li → 網站用 lookup（驗證里名、學校 id；排除未開辦學校；鄰號 key 一律字串；條件改用片段 fr）。"""
    out, problems = {}, []
    raw_by_school = {z['id']: z for z in zones}
    for k, e in by_li.items():
        if k == '_dist_wide':
            continue
        d, li = k.split('|', 1)
        if (d, li) not in vill_keys:
            problems.append(f'lookup key 不在村里界：{k}')
        ent = {'all': [s for s in e.get('all', []) if s not in excluded], 'lin': {}}
        for lin, ss in sorted(e.get('lin', {}).items(), key=lambda x: int(x[0])):
            ss = [s for s in ss if s not in excluded and s not in ent['all']]
            if ss:
                ent['lin'][str(int(lin))] = ss
        if e.get('free'):
            ent['free'] = {s: v for s, v in e['free'].items() if s not in excluded}
        # 片段：只留「不是單一學校、沒條件」的鄰（其餘網頁端當成一個沒條件的片段）
        fr = {}
        for L, lst in (frags.get(k) or {}).items():
            lst = [dict(f, s=[s for s in f['s'] if s not in excluded]) for f in lst]
            lst = [f for f in lst if f['s']]
            allowed = set(ent['all']) | (set(ent['lin'].get(L, [])) if L != 'all' else set())
            if L == 'all':
                allowed |= {s for v in ent['lin'].values() for s in v}
            for f in lst:
                bad = [s for s in f['s'] if s not in allowed]
                if bad:
                    problems.append(f'{k} 第{L}鄰 片段的學校 {bad} 不在查詢表')
            trivial = len(lst) == 1 and not lst[0]['c'] and not lst[0]['f'] and not lst[0]['co']
            if not trivial and lst:
                fr[L] = [frag_json(f) for f in lst]
        if fr:
            ent['fr'] = fr
        for s in set(ent['all']) | {x for v in ent['lin'].values() for x in v}:
            if s not in school_ids:
                problems.append(f'{k} 的學校 {s} 不在學區表學校清單')
        out[k] = ent
    wide = {}
    for d, lst in (by_li.get('_dist_wide') or {}).items():
        wide[d] = []
        for w in lst:
            if w['school'] in excluded:
                continue
            if w['school'] not in school_ids:
                problems.append(f'_dist_wide 的學校 {w["school"]} 不在學區表學校清單')
            z = raw_by_school.get(w['school'])
            raw = ''
            if z:
                raw = next((r.get('raw', '') for r in z['ranges'] if (r.get('whole_dist') or r['li'] == '*') and r['dist'] == d), '')
            o = {'school': w['school'], 'free_zone': bool(w.get('free_zone')), 'co_zone': bool(w.get('co_zone')), 'raw': raw}
            if w['school'] in WIDE_INFO:
                o['info'] = WIDE_INFO[w['school']]
            wide[d].append(o)
    out['_dist_wide'] = wide
    for (d, li) in vill_keys:
        out.setdefault(f'{d}|{li}', {'all': [], 'lin': {}})
    return out, problems


def check_zone_vs_lookup(level, zones, lk):
    """學區表每一筆範圍（學校×里×鄰）都要在查詢表查得到該校；「全里」要在 all 或每個有列的鄰裡。"""
    bad = []
    for z in zones:
        sid = z['id']
        for r in z['ranges']:
            if r.get('whole_dist') or r['li'] == '*':
                if not any(w['school'] == sid for w in lk['_dist_wide'].get(r['dist'], [])):
                    bad.append(f'{level} {sid}：整區 {r["dist"]} 不在 _dist_wide')
                continue
            e = lk.get(f'{r["dist"]}|{r.get("li_norm") or r["li"]}')
            if not e:
                bad.append(f'{level} {sid}：{r["dist"]}{r["li"]} 查詢表沒有這個里')
                continue
            if sid in e['all']:
                continue
            lins = r['lin'] if r['lin'] != 'all' else [int(k) for k in e['lin']]
            miss = [n for n in lins if sid not in e['lin'].get(str(n), [])]
            if r['lin'] == 'all' and not lins:
                miss = ['全里']
            if miss:
                bad.append(f'{level} {sid}：{r["dist"]}{r["li"]} 第 {miss[:6]} 鄰在查詢表找不到這校（原文 {r.get("raw", "")[:40]}）')
    return bad


def schools_at(lk, dist, li, lin, with_wide=True):
    e = lk.get(f'{dist}|{li}')
    s = []
    if e:
        s = list(e['all']) + [x for x in e['lin'].get(str(lin), []) if x not in e['all']]
    if with_wide:
        s += [w['school'] for w in lk['_dist_wide'].get(dist, []) if w['school'] not in s]
    return s


def frags_at(lk, dist, li, lin):
    """某里某鄰的片段（沒列 fr 的學校補一個沒條件的片段）。"""
    e = lk.get(f'{dist}|{li}') or {'all': [], 'lin': {}}
    fr = e.get('fr') or {}
    lst = [dict(f) for f in fr.get(str(lin), []) + fr.get('all', [])]
    have = {s for f in lst for s in f['s']}
    for s in schools_at(lk, dist, li, lin, with_wide=False):
        if s not in have:
            lst.append({'s': [s], 'c': ''})
    return lst


def expected_cells(zones, max_lin):
    """學區表原文逐筆 → 每個「區|里 × 鄰」每校應該有的片段（跟 build_fragments 分開算，當對照組）：
    {(k, 鄰): {校id: {'u': 有「不分路段」的部分, 'c': {條件: 是否共同學區}, 'raw': [原文]}}}。
    規則：本校自己的鄰別條件（base_partial / cond）→ 本校在那一鄰只有那個條件；
          【…共同學區】註記另外帶條件（例：14鄰(軍福十三路以南)）→ 本校不分路段＋「條件＋共同學區」兩段；
          整筆都是共同學區 → 條件照範圍本身。"""
    exp = collections.defaultdict(lambda: collections.defaultdict(lambda: {'u': False, 'c': {}, 'raw': []}))
    for z in zones:
        sid = z['id']
        for r in z['ranges']:
            if r.get('whole_dist') or r['li'] == '*':
                continue
            k = f'{r["dist"]}|{r.get("li_norm") or r["li"]}'
            bp = base_partial(r)
            cond = (r.get('cond') or '').strip()
            co = bool(r.get('co_zone'))
            cl = co_lins_of(r) if co else None
            if r['lin'] == 'all':
                ml = max([(max_lin or {}).get(k) or 0] + [int(n) for n in bp] + [int(n) for n in (cl or {})])
                cover = range(1, ml + 1)
            else:
                cover = r['lin']
            for n in cover:
                L = str(n)
                bc = (bp.get(L) or cond).strip()
                e = exp[(k, L)][sid]
                if r.get('raw') and r['raw'] not in e['raw']:
                    e['raw'].append(r['raw'])

                def add(c, isco):
                    if c:
                        e['c'][c] = e['c'].get(c, False) or isco
                    else:
                        e['u'] = True
                if not co:
                    add(bc, False)
                elif cl is None:
                    add(bc, True)
                elif int(n) in cl:
                    cc = (cl[int(n)] or '').strip()
                    if not cc or cc == bc:
                        add(bc, True)
                    else:
                        add(bc, False)
                        add(cc, True)
                else:
                    add(bc, False)
    return exp


def web_frags(lk, k, L):
    """跟 public/js/sd-lookup.js 的 fragsAt 同一套：fr[鄰] ＋ fr.all，查詢表有列、片段沒列到的學校補一個沒條件的片段。"""
    e = lk.get(k) or {'all': [], 'lin': {}}
    fr = e.get('fr') or {}
    lst = [{'s': list(f['s']), 'c': f.get('c') or '', 'co': bool(f.get('co'))} for f in fr.get(L, []) + fr.get('all', [])]
    have = {s for f in lst for s in f['s']}
    for s in list(e.get('all') or []) + list((e.get('lin') or {}).get(L, [])):
        if s not in have:
            lst.append({'s': [s], 'c': '', 'co': False})
            have.add(s)
    return lst


def check_cond_fragments(level, zones, lk, max_lin):
    """一致性檢查：學區表原文有路段條件的鄰，網頁拿到的片段要忠實反映原文。
    - 原文本校「不分路段」也有 → 一定要有本校不帶條件的片段（例：建功國小水景里全里＋9鄰 74號快速道路以東為共同學區）
    - 原文本校只有帶條件的部分 → 不能多出本校不帶條件的片段（例：旭光國小學田里 12 鄰「312巷105弄以外」）
    - 每個條件都要在、原文寫共同學區的條件片段要標共同學區
    回傳 (錯誤清單, 統計, 原文就沒有不分路段的部分清單)。"""
    exp = expected_cells(zones, max_lin)
    cells = sorted({kl for kl, by in exp.items() if any(v['c'] for v in by.values())},
                   key=lambda x: (x[0], int(x[1])))
    bad, cond_only, n_pairs = [], [], 0
    for k, L in cells:
        by = exp[(k, L)]
        frs = web_frags(lk, k, L)
        out_ids = {s for f in frs for s in f['s']}
        where = f'{level} {k.replace("|", "")} 第{L}鄰'
        if not any(v['u'] for v in by.values()):
            cond_only.append({'cell': where, 'schools': {s: sorted(v['c']) for s, v in by.items()}})
        for sid in sorted(set(by) | out_ids):
            n_pairs += 1
            v = by.get(sid)
            mine = [f for f in frs if sid in f['s']]
            if v is None:
                bad.append(f'{where}：{sid} 學區表這一鄰沒列這校，片段卻有（{[f["c"] for f in mine]}）')
                continue
            out_u = any(not f['c'] for f in mine)
            out_c = {}
            for f in mine:
                if f['c']:
                    out_c[f['c']] = out_c.get(f['c'], False) or (f['co'] and len(f['s']) > 1)
            src = '／'.join(v['raw'])[:120]
            if v['u'] and not out_u:
                bad.append(f'{where}：{sid} 原文有不分路段的部分，片段只剩帶條件的 {sorted(out_c)}（原文 {src}）')
            if not v['u'] and out_u:
                bad.append(f'{where}：{sid} 原文只有 {sorted(v["c"])}，片段卻多了不分路段（原文 {src}）')
            for c, isco in v['c'].items():
                if c not in out_c:
                    bad.append(f'{where}：{sid} 少了條件「{c}」（原文 {src}）')
                elif isco and not out_c[c]:
                    bad.append(f'{where}：{sid} 條件「{c}」原文是共同學區，片段沒標（原文 {src}）')
            for c in out_c:
                if c not in v['c']:
                    bad.append(f'{where}：{sid} 片段多出原文沒有的條件「{c}」（原文 {src}）')
    stats = {'cells': len(cells), 'school_cells': n_pairs, 'cond_only_cells': len(cond_only)}
    return bad, stats, cond_only


# ─────────────────────────── 學區多邊形 ───────────────────────────
def clean_m(g):
    """補掉小洞、丟掉碎片（公尺座標）。"""
    if g is None or g.is_empty:
        return g
    polys = []
    for p in shapely.get_parts(g):
        if p.geom_type != 'Polygon' or p.area < MIN_PART_M2:
            continue
        holes = [r for r in p.interiors if shapely.Polygon(r).area >= MIN_HOLE_M2]
        polys.append(shapely.Polygon(p.exterior, holes))
    if not polys:
        return shapely.Polygon()
    return shapely.make_valid(shapely.MultiPolygon(polys)) if len(polys) > 1 else polys[0]


def label_pieces(lk, pts_xy, pts_meta, city_m, stats):
    """全市門牌 Voronoi：標籤＝門牌里鄰對到的學校組合（不含整區學區）。
    回傳 [(標籤 tuple, 修順後的多邊形)]，彼此不重疊、不留縫（coverage）。"""
    labels = [tuple(sorted(schools_at(lk, d, li, lin, with_wide=False))) for d, li, lin in pts_meta]
    # 同一個位置（1 公尺內）多個門牌 → 一個點，標籤取多數
    key = np.round(pts_xy, 0)
    groups = collections.defaultdict(list)
    for i, k in enumerate(map(tuple, key)):
        groups[k].append(i)
    xy, lab = [], []
    for ids in groups.values():
        c = collections.Counter(labels[i] for i in ids)
        xy.append(pts_xy[ids[0]])
        lab.append(c.most_common(1)[0][0])
    xy = np.array(xy)
    stats['voronoi_points'] = len(xy)
    minx, miny = xy.min(axis=0)
    maxx, maxy = xy.max(axis=0)
    cb = city_m.bounds
    env = shapely.box(min(minx, cb[0]) - 5000, min(miny, cb[1]) - 5000, max(maxx, cb[2]) + 5000, max(maxy, cb[3]) + 5000)
    cells = shapely.get_parts(shapely.voronoi_polygons(shapely.multipoints(xy), extend_to=env, ordered=True))
    if len(cells) != len(xy):
        raise RuntimeError(f'voronoi 產出 {len(cells)} 格，門牌點 {len(xy)}')
    by = collections.defaultdict(list)
    for c, l in zip(cells, lab):
        by[l].append(c)
    labs = sorted(by)
    geoms = [shapely.intersection(shapely.coverage_union_all(np.array(by[l], dtype=object)), city_m) for l in labs]
    geoms = [shapely.make_valid(g) for g in geoms]
    geoms = [shapely.union_all([p for p in shapely.get_parts(g) if p.geom_type == 'Polygon']) for g in geoms]
    geoms = np.array(list(shapely.set_precision(np.array(geoms, dtype=object), 0.01)), dtype=object)
    keep = [i for i, g in enumerate(geoms) if g is not None and not g.is_empty]
    labs, geoms = [labs[i] for i in keep], geoms[keep]
    stats['labels'] = len(labs)
    try:
        if shapely.coverage_is_valid(geoms):
            sm = shapely.coverage_simplify(geoms, VORONOI_SMOOTH_M, simplify_boundary=True)
            if all(g is not None and g.is_valid for g in sm):
                geoms = sm
                stats['smooth'] = 'ok'
            else:
                stats['smooth'] = 'invalid-after（用未修順的邊）'
        else:
            stats['smooth'] = 'coverage-invalid（用未修順的邊）'
    except Exception as ex:   # 修順失敗就用原始鋸齒邊（仍然正確，只是邊比較碎、檔案比較大）
        stats['smooth'] = 'error: ' + str(ex)[:60]
    return [(l, g) for l, g in zip(labs, geoms) if g is not None and not g.is_empty]


def color_schools(feats_m, wide_ids):
    """相鄰（含重疊）的學區不同色；貪婪 DSatur，回傳 {id: 色號}。"""
    ids = [i for i in feats_m if i not in wide_ids and not feats_m[i].is_empty]
    geoms = [feats_m[i] for i in ids]
    tree = shapely.STRtree(geoms)
    nb = {i: set() for i in ids}
    for a, g in enumerate(geoms):
        for b in tree.query(g, predicate='dwithin', distance=5.0):
            if b != a:
                nb[ids[a]].add(ids[b])
                nb[ids[b]].add(ids[a])
    color = {}
    todo = set(ids)
    while todo:
        def sat(i):
            return (len({color[j] for j in nb[i] if j in color}), len(nb[i]), i)
        i = max(todo, key=sat)
        used = {color[j] for j in nb[i] if j in color}
        c = 0
        while c in used:
            c += 1
        color[i] = c
        todo.discard(i)
    return color


def round_geom(g):
    gj = mapping(g)

    def rnd(c):
        if isinstance(c, (list, tuple)) and c and isinstance(c[0], (int, float)):
            return [round(c[0], 5), round(c[1], 5)]
        return [rnd(x) for x in c]
    return {'type': gj['type'], 'coordinates': rnd(gj['coordinates'])}


def school_relations(lk):
    """每校：shared＝學區表寫明「共同學區」的其他學校；split＝同一鄰依路段／條件拆給的其他學校
    （地圖上這兩種都會重疊，因為沒有官方鄰界圖可以再切）。"""
    shared, split, choice = collections.defaultdict(set), collections.defaultdict(set), collections.defaultdict(set)
    for k, e in lk.items():
        if k.startswith('_'):
            continue
        d, li = k.split('|', 1)
        lins = set(e['lin']) | set((e.get('fr') or {}).keys()) | {'0'}
        for L in lins:
            if L == 'all':
                continue
            frs = frags_at(lk, d, li, L)
            ss = {s for f in frs for s in f['s']}
            in_co = {s for f in frs if f.get('co') and not f.get('q') for s in f['s']}
            in_ch = {s for f in frs if f.get('co') and f.get('q') for s in f['s']}
            for a in ss:
                for b in ss - {a}:
                    cond_split = any(f.get('c') and a in f['s'] and b not in f['s'] for f in frs) or \
                        any(f.get('c') and b in f['s'] and a not in f['s'] for f in frs)
                    both_co = any(f.get('co') and not f.get('q') and a in f['s'] and b in f['s'] for f in frs) or (a in in_co and b in in_co)
                    both_ch = a in in_ch and b in in_ch and not both_co
                    if cond_split:
                        split[a].add(b)          # 依路段／條件拆給兩校
                    if both_co:
                        shared[a].add(b)         # 學區表寫明共同學區（同一段，或兩校那一列各自寫了共同學區）
                    if both_ch:
                        choice[a].add(b)         # 學區表註記「可選…就讀」（原文沒寫共同學區）
                    if not cond_split and not both_co and not both_ch:
                        split[a].add(b)          # 同一鄰列在兩校、沒寫共同學區也沒條件 → 請用地址查詢
    return shared, split, choice


def build_zones(level_key, level, zones, lk, pieces, vill_m, geo_by_id, stats):
    by_school = collections.defaultdict(list)
    for lab, g in pieces:
        for s in lab:
            by_school[s].append(g)
    shared, split, choice = school_relations(lk)
    # 整區共同 / 自由學區：用村里界拼整區
    wide_by_school = collections.defaultdict(list)
    wide_raw = {}
    for d, lst in lk['_dist_wide'].items():
        for w in lst:
            wide_by_school[w['school']].append(d)
            wide_raw[w['school']] = w.get('raw', '')
    feats_m = {}
    for z in zones:
        sid = z['id']
        gs = list(by_school.get(sid, []))
        if sid in wide_by_school:
            gs += [g for (d, li), g in vill_m.items() if d in wide_by_school[sid]]
            feats_m[sid] = clean_m(shapely.union_all(gs)) if gs else shapely.Polygon()
        else:
            feats_m[sid] = clean_m(shapely.coverage_union_all(np.array(gs, dtype=object))) if gs else shapely.Polygon()
    colors = color_schools(feats_m, set(wide_by_school))
    n_li = collections.Counter()
    approx = set()
    for k, e in lk.items():
        if k.startswith('_'):
            continue
        for s in set(e['all']) | {x for v in e['lin'].values() for x in v}:
            n_li[s] += 1
        if e['lin']:
            for v in e['lin'].values():
                approx |= set(v)
    free_s = {s for k, e in lk.items() if not k.startswith('_') for s in (e.get('free') or {})}
    feats, empty = [], []
    for z in zones:
        sid = z['id']
        g = feats_m[sid]
        if g.is_empty:
            empty.append(sid)
            continue
        gd = shapely.make_valid(shapely.set_precision(to_deg(g), 1e-5))
        gd = shapely.union_all([p for p in shapely.get_parts(gd) if p.geom_type == 'Polygon'])
        geo = geo_by_id[sid]
        props = {
            'school': sid, 'name': z['school'], 'level': level, 'dist': geo['dist'],
            'n_li': n_li[sid], 'shared': sorted(shared[sid]), 'approx': sid in approx,
        }
        if split[sid]:
            props['split'] = sorted(split[sid])
        if choice[sid]:
            props['choice'] = sorted(choice[sid])
        if sid in wide_by_school:
            props['wide'] = True
            props['wide_dists'] = wide_by_school[sid]
            props['wide_raw'] = wide_raw.get(sid, '')
            props['approx'] = False
        if sid in free_s or any(w.get('free_zone') for d in wide_by_school.get(sid, []) for w in lk['_dist_wide'][d] if w['school'] == sid):
            props['free'] = True
        if sid in colors:
            props['ci'] = colors[sid]
        feats.append({'type': 'Feature', 'properties': props, 'geometry': round_geom(gd)})
    stats[f'{level_key}_empty_zone'] = empty
    stats[f'{level_key}_colors'] = max(colors.values()) + 1 if colors else 0
    return {'type': 'FeatureCollection', 'meta': {'level': level, 'year': YEAR}, 'features': feats}


def check_polygons(level, lk, fc, pts_lnglat, pts_meta):
    """每個門牌點落在哪些學校的多邊形裡；落在「只有別校」的比例（交界修順造成的誤差）。"""
    X = np.array([p[0] for p in pts_lnglat])
    Y = np.array([p[1] for p in pts_lnglat])
    hits = [set() for _ in range(len(X))]
    for f in fc['features']:
        if f['properties'].get('wide'):
            continue
        g = shape(f['geometry'])
        shapely.prepare(g)
        b = g.bounds
        sel = np.nonzero((X >= b[0]) & (X <= b[2]) & (Y >= b[1]) & (Y <= b[3]))[0]
        if not len(sel):
            continue
        for i in sel[shapely.contains_xy(g, X[sel], Y[sel])]:
            hits[i].add(f['properties']['school'])
    wrong = collections.Counter()
    n_wrong = n_none = 0
    for i, (d, li, lin) in enumerate(pts_meta):
        truth = set(schools_at(lk, d, li, lin, with_wide=False))
        if not truth:
            continue
        if not hits[i]:
            n_none += 1
        elif not (hits[i] & truth):
            n_wrong += 1
            wrong[f'{d}{li}{lin}鄰'] += 1
    return {'points': len(X), 'only_other_school': n_wrong, 'pct': round(n_wrong * 100 / len(X), 4),
            'no_polygon': n_none, 'top': wrong.most_common(15)}


# ─────────────────────────── 學校資訊 ───────────────────────────
def compress_lins(ns):
    ns = sorted(set(ns))
    out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(f'{ns[i]}-{ns[j]}' if j > i + 1 else (f'{ns[i]}、{ns[j]}' if j == i + 1 else f'{ns[i]}'))
        i = j + 1
    return '、'.join(out)


def lins_with_conds(m):
    """{鄰: 條件} → 「第3-6、12、17鄰、第14鄰（大慶街二段以南）」：沒條件的鄰先列（壓成區間），帶條件的鄰各自寫出條件
    （跟學區表原文「第3-6、12…鄰及14鄰(大慶街二段以南)」同一個順序）。"""
    plain = [n for n in m if not m[n]]
    out = [f'第{compress_lins(plain)}鄰'] if plain else []
    out += [f'第{n}鄰（{m[n]}）' for n in sorted(m) if m[n]]
    return '、'.join(out)


def range_desc(r, co_lins=None):
    """學校彈窗「學區範圍（里、鄰）」的一列。co_lins：{鄰: 共同學區那一段另外的路段條件}（只有部分鄰、
    或共同學區只限某路段時才有），例：東山高中國中部和平里「全里【第16、33、35鄰、第14鄰（軍福十三路以南）、
    第36鄰（軍福十三路以南）為共同學區】」。本校自己的鄰別條件用 base_partial（註記裡共同學區的條件不算本校的）。"""
    lins = r['lin']
    part = {int(k): v for k, v in base_partial(r).items()}
    rule = r.get('lin_rule') or {}
    if lins == 'all':
        d = '全里'
        if part:     # 學田里(全里)(12鄰:312巷105弄以外)
            d += '（' + '、'.join(f'第{n}鄰：{part[n]}' for n in sorted(part)) + '）'
    elif rule.get('except'):
        d = f'全里（第{compress_lins(rule["except"])}鄰除外）'
        for n in sorted(part):
            d += f'、第{n}鄰（{part[n]}）'
    else:
        d = lins_with_conds({n: part.get(n, '') for n in lins})
    if r.get('cond'):
        d += f'（{r["cond"]}）'
    fz = r.get('free_zone')
    if fz is True:
        d += '【自由學區】'
    elif isinstance(fz, list) and fz:
        d += f'【第{compress_lins(fz)}鄰為自由學區】'
    if r.get('choice'):
        # 「頭汴11、12鄰可選坪林國小就讀」：原文沒寫共同學區 → 照原文引用，不寫成共同學區
        d += f'（學區表註記：{r["choice"]}）'
    elif r.get('co_zone'):
        # 原文「其中第12、26鄰為北新國中共同學區」：只有部分鄰是共同學區時寫出是哪幾鄰；
        # 共同學區只限某路段（14鄰(軍福十三路以南)）時，條件寫在那一鄰後面，不能寫成整鄰共同學區
        d += f'【{lins_with_conds(co_lins)}為共同學區】' if co_lins else '【共同學區】'
    return d


def co_partners(lk, dist, li, r, sid, max_lin=None):
    """只有原文寫明「共同學區」的那一列才掛共同學區的學校：同一鄰、同一段共同學區片段裡的其他學校；
    片段裡只有本校（各校那一列寫法不同）時，改列同一鄰的其他學校。
    回傳 (共同學區學校, {鄰: 共同學區那一段另外的路段條件} 或 None)：
    只有部分鄰是共同學區、或某鄰的共同學區只限某路段（本校那一鄰不帶條件）時才回 dict。"""
    if not r.get('co_zone') or r.get('choice'):
        return [], None
    ml = (max_lin or {}).get(f'{dist}|{li}') or 0
    lins = [str(n) for n in r['lin']] if r['lin'] != 'all' else ([str(n) for n in range(1, ml + 1)] or ['0'])
    bp = base_partial(r)
    out, co_map = set(), {}
    for L in lins:
        frs = frags_at(lk, dist, li, L)
        mine = [f for f in frs if f.get('co') and sid in f['s']]
        here = set()
        for f in mine:
            if len(f['s']) > 1:
                here |= set(f['s']) - {sid}
            else:
                here |= {s for g in frs for s in g['s']} - {sid}
        if here:
            base_c = (bp.get(L) or r.get('cond') or '').strip()
            co_map[int(L)] = '；'.join(sorted({f['c'] for f in mine if f['c'] and f['c'] != base_c}))
        out |= here
    part = co_map if (co_map and lins != ['0'] and (len(co_map) < len(lins) or any(co_map.values()))) else None
    return sorted(out), part


def build_school_info(level, zones, lk, geo_by_id, moe_phones, new_codes, max_lin=None):
    info = {}
    for z in zones:
        sid = z['id']
        g = geo_by_id[sid]
        lis, raws = [], []
        for r in z['ranges']:
            if r.get('whole_dist') or r['li'] == '*':
                lis.append([r['dist'], '*', '全區' + ('【自由學區】' if r.get('free_zone') else '【共同學區】' if r.get('co_zone') else '')])
                raws.append(r.get('raw', ''))
                continue
            li = r.get('li_norm') or r['li']
            partners, co_lins = co_partners(lk, r['dist'], li, r, sid, max_lin)
            lis.append([r['dist'], li, range_desc(r, co_lins)] + partners)
            raws.append(r.get('raw', ''))
        phone = norm_phone(z.get('phone')) or norm_phone(moe_phones.get(g.get('moe_code') or '', ''))
        rec = {
            'name': z['school'], 'level': level, 'dist': g['dist'], 'addr': re.sub(r'^臺中市', '', g['addr']),
            'lat': round(g['lat'], 6), 'lng': round(g['lng'], 6),
            'phone': phone, 'lis': lis, 'raw': raws,
        }
        if (g.get('moe_code') or '') in new_codes:
            rec['status'] = '115 學年度新設校'
            rec['status_src'] = BASIC_CHANGE_URL
        if g.get('geo_precision') and g['geo_precision'] != 'exact':
            rec['geo'] = g['geo_precision']
        if z.get('notes'):
            rec['notes'] = z['notes']
        info[sid] = rec
    return info


def split_tail_notes(desc):
    """「第1-6鄰、第7鄰（民生路【140號】）【第3鄰為共同學區】」→ ('第1-6鄰、第7鄰（民生路【140號】）', ['【第3鄰為共同學區】'])：
    只切掉 range_desc 接在最後面的【自由學區】【…共同學區】註記（條件文字裡本來就有的【…】不動）。"""
    tails = []
    while desc.endswith('】'):
        depth, i = 0, len(desc) - 1
        while i >= 0:
            ch = desc[i]
            if ch == '】':
                depth += 1
            elif ch == '【':
                depth -= 1
                if depth == 0:
                    break
            i -= 1
        if i < 0:
            break
        tails.insert(0, desc[i:])
        desc = desc[:i]
    return desc, tails


def check_lis_conds(level, zones, info):
    """一致性檢查：學校彈窗的學區範圍描述，路段條件要掛在正確的地方。
    - 本校自己的鄰別條件（base_partial、cond）要寫在本校範圍裡
    - 只屬共同學區那一段的條件（例：四育國中樹德里 14 鄰「大慶街二段以南」）不能寫成本校範圍的條件，
      要寫在【…為共同學區】裡那一鄰的後面（例：東山高中國中部和平里 14 鄰「軍福十三路以南」）
    回傳 (錯誤清單, 檢查的條件數)。"""
    bad, n = [], 0
    for z in zones:
        rec = info.get(z['id']) or {}
        for r, row in zip(z['ranges'], rec.get('lis') or []):
            if r.get('whole_dist') or r['li'] == '*':
                continue
            desc = row[2]
            head, tails = split_tail_notes(desc)
            co_seg = next((t[1:-len('為共同學區】')] for t in tails if t.endswith('為共同學區】')), '')
            where = f'{level} {z["id"]} {r["dist"]}{row[1]}「{desc}」'
            bp = base_partial(r)
            cond = (r.get('cond') or '').strip()
            for L, c in bp.items():
                n += 1
                if f'第{L}鄰（{c}）' not in head and f'第{L}鄰：{c}' not in head:
                    bad.append(f'{where}：少了本校第{L}鄰的條件「{c}」（原文 {r.get("raw", "")[:80]}）')
            for L, c in (r.get('partial') or {}).items():
                c = (c or '').strip()
                if str(L) not in bp:
                    n += 1
                    if c in head:
                        bad.append(f'{where}：第{L}鄰「{c}」原文只屬共同學區那一段，卻寫成本校範圍的條件（原文 {r.get("raw", "")[:80]}）')
            cl = co_lins_of(r) if r.get('co_zone') else None
            for L, cc in (cl or {}).items():
                cc = (cc or '').strip()
                if cc and cc != (bp.get(str(L)) or cond):
                    n += 1
                    if f'第{L}鄰（{cc}）' not in co_seg:
                        bad.append(f'{where}：共同學區只限第{L}鄰「{cc}」，描述沒寫出來（原文 {r.get("raw", "")[:80]}）')
            if cond:
                n += 1
                if f'（{cond}）' not in desc:
                    bad.append(f'{where}：少了條件「{cond}」（原文 {r.get("raw", "")[:80]}）')
    return bad, n


# ─────────────────────────── 主程式 ───────────────────────────
def abort(fail, report, path):
    """一致性檢查沒過：寫報告、不動網站資料檔、exit 1。"""
    report['fail'] = fail
    os.makedirs(os.path.dirname(path), exist_ok=True)
    jdump(report, path)
    log('\n❌ 一致性檢查沒過（網站資料檔未更新）：')
    for f in fail:
        log('  - ' + f)
    log(f'報告：{path}')
    sys.exit(1)


def main():
    ap = argparse.ArgumentParser(description='台中市國中小學區查詢地圖資料建置')
    ap.add_argument('--sd', default=DEFAULT_SD, help='上游學區表資料夾（zones_*.json / by_li_*.json）')
    ap.add_argument('--geo', default=DEFAULT_GEO, help='上游地理資料夾（villages_tc.geojson / schools.json / raw/moe115）')
    ap.add_argument('--addr', default=DEFAULT_ADDR, help='門牌庫資料夾（scripts/build-tc-addr.py 的輸出）')
    ap.add_argument('--out', default=OUT_DIR)
    ap.add_argument('--report', default=DEFAULT_REPORT, help='建置報告 JSON（不進 repo）')
    a = ap.parse_args()
    t0 = time.time()
    fail = []
    report = {'built_at': time.strftime('%Y-%m-%d %H:%M:%S'), 'year': YEAR}

    log('[1] 載入上游資料')
    I = load_inputs(a.sd, a.geo)
    vill = I['villages']
    vill_keys = {(f['properties']['dist'], f['properties']['li']) for f in vill['features']}
    li_dists = collections.defaultdict(set)
    for d, li in vill_keys:
        li_dists[li].add(d)
    zones = {'gs': I['zones_gs'], 'jh': I['zones_jh']}
    moe_phones = load_moe_phones(a.geo)
    new_codes = load_new_schools(a.geo)
    log(f'  教育部名錄電話 {len(moe_phones)} 校；115 新設校代碼 {sorted(new_codes)}')

    log('[2] 排除籌備中／預定學校')
    excluded = {}
    for lv, _ in LEVELS:
        keep = []
        for z in zones[lv]:
            if NOT_OPEN_RE.search(z['school']) or NOT_OPEN_RE.search(z.get('status') or ''):
                excluded[z['id']] = z.get('status') or z['school']
            else:
                keep.append(z)
        zones[lv] = keep
    report['excluded_not_open'] = excluded
    log(f'  排除 {len(excluded)} 校')

    # 「※【共同學區可選擇就讀頭家國小或僑忠國小】」這種 ※ 註記是整校的說明，上游把它接在最後一個里後面
    # （頭家國小的家興里原文是「家興里(全里)」）→ ※ 前面沒寫共同學區的，不算這個里的共同學區
    star_fix = []
    for lv, _ in LEVELS:
        for z in zones[lv]:
            for r in z['ranges']:
                raw = r.get('raw') or ''
                if r.get('co_zone') and '※' in raw and '共同學區' not in raw.split('※', 1)[0]:
                    r['co_zone'] = False
                    r['ann'] = ''
                    star_fix.append(f'{z["id"]} {r["dist"]}{r["li"]}：{raw}')
    report['star_note_fixed'] = star_fix
    log(f'  ※ 註記不算該里共同學區：{star_fix}')

    # 「頭汴11、12鄰可選坪林國小就讀」：上游把「可選…就讀」當成共同學區（co_zone），但原文沒寫「共同學區」
    # → 標 choice（註記原文），片段帶 q、學校彈窗照原文引用；網頁顯示「學區表註記：『…』」，⛔ 不寫成共同學區
    choice_fix = []
    for lv, _ in LEVELS:
        for z in zones[lv]:
            for r in z['ranges']:
                txt = (r.get('ann') or '') + (r.get('raw') or '')
                if r.get('co_zone') and '可選' in txt and '共同' not in txt:
                    qs = re.findall(r'【([^【】]*可選[^【】]*)】', txt)
                    r['choice'] = '；'.join(dict.fromkeys(qs)) if qs else re.sub(r'[【】]', '', r.get('ann') or '')
                    choice_fix.append(f'{z["id"]} {r["dist"]}{r["li"]}：{r["choice"]}')
    report['choice_note'] = choice_fix
    log(f'  「可選…就讀」註記（不寫成共同學區）：{choice_fix}')

    # 同一鄰依門牌拆給兩校（清水區南社里 9、27 鄰；裕嘉里 1 鄰）：上游的條件後面接了「（原文：另一校那一半…）」
    # 「（備註：…）」→ 結果卡會變成「中華路雙號及南華路（原文：…為建國國小）：清水國小」，同一行出現兩所學校。
    # 條件只留本校那一段（完整原文仍在學校 raw／notes）
    tail_re = re.compile(r'（(?:原文|備註)：.*）$')
    cond_fix = []
    for lv, _ in LEVELS:
        for z in zones[lv]:
            for r in z['ranges']:
                part = r.get('partial') or {}
                for L in list(part):
                    c = part[L] or ''
                    if tail_re.search(c):
                        part[L] = tail_re.sub('', c).strip()
                        cond_fix.append(f'{z["id"]} {r["dist"]}{r["li"]} 第{L}鄰：{c} → {part[L]}')
                if r.get('cond') and tail_re.search(r['cond']):
                    c = r['cond']
                    r['cond'] = tail_re.sub('', c).strip()
                    cond_fix.append(f'{z["id"]} {r["dist"]}{r["li"]}：{c} → {r["cond"]}')
    report['cond_tail_fixed'] = cond_fix
    log(f'  條件拿掉接在後面的另一校原文：{len(cond_fix)} 筆')

    log('[3] 學區表里名 × 村里界')
    for lv, level in LEVELS:
        fixed, bad = check_range_villages(level, zones[lv], vill_keys, li_dists)
        report[f'{lv}_village_fixed'] = fixed
        report[f'{lv}_village_unmatched'] = bad
        log(f'  {level}：對照表修正 {len(fixed)} 筆，對不上 {len(bad)} 筆')
        fail += [f'學區表里名對不上村里界：{b}' for b in bad]

    log('[4] 學校座標')
    geo_by = {}
    for lv, level in LEVELS:
        m, errs = match_schools(level, zones[lv], I['schools'], li_dists)
        geo_by[lv] = m
        for sid, g in m.items():
            if g.get('lat') is None or g.get('lng') is None:
                errs.append(f'{level} {sid} 沒有座標')
        fail += errs
        log(f'  {level}：{len(m)}/{len(zones[lv])} 校有座標' + (f'；問題 {errs}' if errs else ''))
    if fail:   # 後面的步驟都要用到每校座標，這裡就先停
        abort(fail, report, a.report)

    log('[5] 查詢表（含學區表條件片段）')
    lookups = {}
    for lv, level in LEVELS:
        ids = {z['id'] for z in zones[lv]}
        fixlog = []
        frags = build_fragments(zones[lv], fixlog)
        report[f'{lv}_co_fixes'] = fixlog
        log(f'  {level}：共同學區補強 {len(fixlog)} 筆（明細見報告 {lv}_co_fixes）')
        lk, probs = build_lookup(I[f'by_li_{lv}'], frags, zones[lv], vill_keys, ids, set(excluded))
        co_clean, co_all_single = clean_single_co(lk)
        report[f'{lv}_co_single_cleaned'] = co_clean
        report[f'{lv}_co_all_frag_single_lins'] = co_all_single
        log(f'  {level}：只有一校的鄰拿掉共同學區標記 {len(co_clean)} 個鄰；'
            f'「全里」片段標共同學區、但有些鄰只有這一校的里 {len(co_all_single)} 個（網頁端只有一校時不看 co，燈號不受影響）')
        cbad1 = check_single_co(lk)
        report[f'{lv}_co_single_check'] = cbad1
        fail += cbad1
        lookups[lv] = lk
        fail += probs
        used = {s for k, e in lk.items() if not k.startswith('_') for s in set(e['all']) | {x for v in e['lin'].values() for x in v}}
        used |= {w['school'] for v in lk['_dist_wide'].values() for w in v}
        nozone = sorted(ids - used)
        if nozone:
            fail.append(f'{level} 有學校在查詢表裡完全沒出現：{nozone}')
        zbad = check_zone_vs_lookup(level, zones[lv], lk)
        report[f'{lv}_zone_vs_lookup'] = zbad
        fail += zbad
        n_fr = sum(len(e.get('fr') or {}) for k, e in lk.items() if not k.startswith('_'))
        log(f'  {level}：{len(lk) - 1} 里；有條件／共同學區片段的鄰 {n_fr}；問題 {len(probs)}；學區表範圍 vs 查詢表不一致 {len(zbad)}')
        # 原文有路段條件的鄰：片段要同時有「不分路段」和「帶條件」（原文本校就只有帶條件的部分時，列進報告 cond_only）
        cbad, cst, conly = check_cond_fragments(level, zones[lv], lk, I['max_lin'])
        report[f'{lv}_cond_check'] = {'stats': cst, 'bad': cbad, 'cond_only': conly}
        fail += cbad
        log(f'  {level}：路段條件檢查 {cst["cells"]} 個鄰、{cst["school_cells"]} 個「學校×鄰」，不符 {len(cbad)}；'
            f'原文就沒有不分路段部分的鄰 {cst["cond_only_cells"]}（明細見報告 {lv}_cond_check.cond_only）')

    log('[6] 讀門牌庫（涵蓋率檢查＋畫學區）')
    pts, cnt, alt_cnt, addr_ver, n_addr = load_addr_points(a.addr)
    naddr = len(pts)
    log(f'  門牌庫 {addr_ver}：{naddr} 個門牌，{len(cnt)} 個區|里|鄰；同一門牌也登記在別的里鄰 {sum(alt_cnt.values())} 筆')
    cov = {}
    for lv, level in LEVELS:
        miss = collections.Counter()
        miss_nowide = 0
        for (d, li, lin), n in cnt.items():
            if not schools_at(lookups[lv], d, li, lin):
                miss[f'{d}|{li}|{lin}'] += n
            if not schools_at(lookups[lv], d, li, lin, with_wide=False):
                miss_nowide += n
        alt_miss = {f'{d}|{li}|{lin}': n for (d, li, lin), n in alt_cnt.items() if not schools_at(lookups[lv], d, li, lin)}
        tot = sum(cnt.values())
        mt = sum(miss.values())
        cov[lv] = {
            'addr_total': tot, 'addr_unmatched': mt, 'unmatched_pct': round(mt * 100 / tot, 4),
            'addr_only_via_dist_wide': miss_nowide - mt,
            'keys_total': len(cnt), 'keys_unmatched': len(miss),
            'unmatched': [{'key': k, 'addr': n} for k, n in miss.most_common()],
            'alt_unmatched': alt_miss,
        }
        log(f'  {level}：對不到學校的門牌 {mt}/{tot}（{mt * 100 / tot:.3f}%），{len(miss)} 個區|里|鄰；'
            f'只靠整區學區才對到 {miss_nowide - mt}')
        if mt * 100 / tot > 0.5:
            fail.append(f'{level} 對不到學校的門牌比例 {mt * 100 / tot:.3f}% 超過 0.5%，學區表或里名對照可能有問題')
    report['coverage'] = cov

    log('[7] 學區多邊形（全市門牌 Voronoi）')
    vill_m = {}
    for f in vill['features']:
        p = f['properties']
        vill_m[(p['dist'], p['li'])] = shapely.make_valid(to_m(shape(f['geometry'])))
    city_m = shapely.make_valid(shapely.union_all(list(vill_m.values())))
    lnglat = np.array([(p[0], p[1]) for p in pts])
    X, Y = T_FWD.transform(lnglat[:, 0], lnglat[:, 1])
    pts_xy = np.column_stack([X, Y])
    pts_meta = [(p[2], p[3], p[4]) for p in pts]
    zones_out, stats = {}, {}
    for lv, level in LEVELS:
        st = {}
        pieces = label_pieces(lookups[lv], pts_xy, pts_meta, city_m, st)
        zones_out[lv] = build_zones(lv, level, zones[lv], lookups[lv], pieces, vill_m, geo_by[lv], st)
        pc = check_polygons(level, lookups[lv], zones_out[lv], pts, pts_meta)
        st['point_check'] = pc
        stats[lv] = st
        log(f'  {level}：{len(zones_out[lv]["features"])} 校，{st["voronoi_points"]} 個門牌點、{st["labels"]} 種學校組合，'
            f'修順 {st.get("smooth")}，配色 {st[lv + "_colors"]} 色，無範圍 {st[lv + "_empty_zone"]}；'
            f'門牌點只落在別校範圍 {pc["only_other_school"]}（{pc["pct"]}%）')
        if pc['pct'] > 0.1:
            fail.append(f'{level} 門牌點只落在別校多邊形的比例 {pc["pct"]}% 超過 0.1%')
    report['zones'] = stats

    log('[8] 學校資訊')
    info = {lv: build_school_info(level, zones[lv], lookups[lv], geo_by[lv], moe_phones, new_codes, I['max_lin']) for lv, level in LEVELS}
    for lv, level in LEVELS:
        lbad, ln = check_lis_conds(level, zones[lv], info[lv])
        report[f'{lv}_lis_cond_check'] = {'checked': ln, 'bad': lbad}
        fail += lbad
        log(f'  {level}：學區範圍描述的路段條件 {ln} 處，不符 {len(lbad)}')
    nophone = [f'{lv} {sid}' for lv in info for sid, v in info[lv].items() if not v['phone']]
    report['no_phone'] = nophone
    log(f'  沒有電話的學校：{nophone or "無"}')

    # meta
    sdr, gr = I['sd_report'], I['geo_report']
    today = datetime.date.today().isoformat()
    gs_tab = {'name': '臺中市政府教育局「115學年度本市國民小學學區表」（中區區公所轉公告，附學區調整說明表）',
              'url': sdr['sources']['gs115.pdf']['url'], 'page': sdr['sources']['gs115.pdf']['page'],
              'version': '115學年度（2026-02-24 公告）', 'kind': 'official'}
    jh_tab = {'name': '臺中市政府教育局「臺中市115學年度國中學區表」（西屯區公所公告）',
              'url': sdr['sources']['jh115.odt']['url'], 'page': sdr['sources']['jh115.odt']['page'],
              'version': '115學年度（2026-02-11 公告）', 'kind': 'official'}
    meta = {
        'title': '臺中市國民中小學學區', 'year': YEAR, 'yearNote': '2026-08 起適用（115學年度）',
        'built': today,
        'counts': {
            'gs': len(zones['gs']), 'jh': len(zones['jh']), 'li': len(vill_keys), 'addr': naddr,
            'gs_zone_features': len(zones_out['gs']['features']), 'jh_zone_features': len(zones_out['jh']['features']),
        },
        'officialQuery': OFFICIAL_QUERY,
        'tables': [gs_tab, jh_tab],
        'sources': [
            gs_tab, jh_tab,
            {'name': '臺中市政府數位發展局「臺中市空間資訊建物及門牌號碼位置新版本資料」',
             'url': 'https://opendata.taichung.gov.tw/search/50168dd2-4239-4f3e-89cf-f706617f0436',
             'version': f'{addr_ver or "115年1月"}版', 'kind': 'opendata', 'org': '臺中市政府數位發展局', 'year': '2026'},
            {'name': '內政部國土測繪中心「村里界圖(TWD97經緯度)」', 'url': 'https://data.gov.tw/dataset/7438',
             'version': gr['villages']['version'], 'kind': 'opendata', 'org': '內政部國土測繪中心', 'year': '2026'},
            {'name': '教育部統計處「115學年度各級學校名錄」（國民小學、國民中學、附設國中小部；校址、電話）',
             'url': MOE_LIST_URL, 'page': 'https://depart.moe.edu.tw/ed4500/News_Content.aspx?n=63F5AB3D02A8BBAC&sms=1FF9979D10DBF9F3&s=52D3798D5069E231',
             'version': '115學年度', 'kind': 'other'},
            {'name': '教育部統計處「115學年度國中小異動一覽表」（新設校）', 'url': BASIC_CHANGE_URL,
             'version': '115.08.03 公告、115.08.27 更新', 'kind': 'other'},
        ],
        'license': '政府資料開放授權條款－第1版',
        'licenseUrl': 'https://data.gov.tw/license',
        'attribution': '此開放資料依政府資料開放授權條款 (Open Government Data License) 進行公眾釋出，使用者於遵守本條款各項規定之前提下，得利用之。',
        'coverage': {lv: {k: v for k, v in cov[lv].items() if k not in ('unmatched', 'alt_unmatched')} for lv in cov},
        'polygonCheck': {lv: {k: v for k, v in stats[lv]['point_check'].items() if k != 'top'} for lv in stats},
        'notes': [
            '學區以教育局公告的學區表為準；本工具由門牌資料推算里鄰，僅供參考，入學前請向學校或教育局確認。',
            '地圖上的學區交界是依門牌位置推估（全市門牌 Voronoi）的，交界附近以地址查詢為準。',
            '「共同學區」「自由學區」照學區表原文標示；實際怎麼分發、能不能選，以學校說明為準。',
            '國立中科實驗高級中學國小部：學區表寫「大雅區及西屯區皆為共同學區」；該校另外辦理招生。',
            '和平國中：國中學區表在和平國中那一列，行政區寫「和平區」，里鄰寫【自由學區】。',
        ],
    }

    log('[9] 寫檔')
    os.makedirs(a.out, exist_ok=True)
    outs = {
        'zones_gs.json': zones_out['gs'], 'zones_jh.json': zones_out['jh'],
        'lookup_gs.json': lookups['gs'], 'lookup_jh.json': lookups['jh'],
        'school_info.json': info, 'meta.json': meta,
    }
    tmp = {}
    for fn, obj in outs.items():
        p = os.path.join(a.out, fn + '.tmp')
        jdump(obj, p)
        tmp[fn] = p
    zsize = os.path.getsize(tmp['zones_gs.json']) + os.path.getsize(tmp['zones_jh.json'])
    if zsize > 3.2e6:
        fail.append(f'zones 兩檔合計 {zsize / 1e6:.2f}MB 超過 3MB 目標')
    for lv, _ in LEVELS:
        bad_geo = [f['properties']['school'] for f in zones_out[lv]['features']
                   if not shape(f['geometry']).is_valid]
        if bad_geo:
            fail.append(f'{lv} 多邊形不合法：{bad_geo}')
    if fail:
        for p in tmp.values():
            os.remove(p)
        abort(fail, report, a.report)
    report['fail'] = []
    report['sizes'] = {fn: os.path.getsize(p) for fn, p in tmp.items()}
    os.makedirs(os.path.dirname(a.report), exist_ok=True)
    jdump(report, a.report)
    sizes, unchanged = {}, []
    for fn, p in tmp.items():
        dst = os.path.join(a.out, fn)
        if os.path.exists(dst) and os.path.getsize(dst) == os.path.getsize(p):
            with open(dst, 'rb') as f1, open(p, 'rb') as f2:
                same = f1.read() == f2.read()
            if same:              # 內容一字不差 → 不換檔（dev server 正在送這個檔時換不掉，也免得觸發重新載入）
                os.remove(p)
                sizes[fn] = os.path.getsize(dst)
                unchanged.append(fn)
                continue
        # Windows：本機 dev server 開著這個檔時 os.replace 會 PermissionError（瀏覽器中途斷線時 dev server
        # 送檔用的讀取串流不會關，檔案會一直被占住）→ 稍等重試；還是不行就改成原地覆寫（占用的程式有共用寫入）
        for i in range(10):
            try:
                os.replace(p, dst)
                break
            except PermissionError:
                if i < 9:
                    time.sleep(0.5)
                    continue
                with open(p, 'rb') as fh:
                    data = fh.read()
                with open(dst, 'wb') as fh:
                    fh.write(data)
                    fh.flush()
                    os.fsync(fh.fileno())
                with open(dst, 'rb') as fh:
                    if fh.read() != data:
                        raise SystemExit(f'{fn} 原地覆寫後內容不符，請重跑')
                os.remove(p)
                log(f'  {fn}：檔案被別的程式開著，改用原地覆寫')
        sizes[fn] = os.path.getsize(dst)
    # 舊版副檔名（.geojson）留著會被誤用 → 刪掉
    for old in ('zones_gs.geojson', 'zones_jh.geojson'):
        op = os.path.join(a.out, old)
        if os.path.exists(op):
            os.remove(op)
    log('\n✅ 完成（%.1f 秒）' % (time.time() - t0))
    for fn, sz in sizes.items():
        log(f'  {fn:22s} {sz / 1024:9.1f} KB' + ('（內容沒變，沿用原檔）' if fn in unchanged else ''))
    log(f'  zones 合計 {zsize / 1e6:.2f} MB；報告 {a.report}')


if __name__ == '__main__':
    main()
